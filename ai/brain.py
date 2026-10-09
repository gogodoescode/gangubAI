"""GangubAI brain: LangGraph agent (Groq), RAG over the slides, and the tools it can call.

Tools: retrieve_context, move_robot, set_wander_mode, timer, pomodoro.
`move()` / `set_wander()` are also used directly by the voice loop's keyword shortcut.
"""

from __future__ import annotations

import os
import re
import subprocess
from typing import Annotated, Literal, TypedDict
from uuid import uuid4

from langchain_core.messages import AIMessage, BaseMessage, HumanMessage, SystemMessage, ToolMessage
from langchain_core.tools import tool
from langchain_groq import ChatGroq
from langgraph.graph import START, StateGraph
from langgraph.graph.message import add_messages
from langgraph.prebuilt import ToolNode, tools_condition
from pydantic import BaseModel, Field

from ai import config
from ai.face import send_to_face


# ═══════════════════════════════════════════════════════════════════
# Robot commands (ROS 2 via the `ros2` CLI, no rclpy node needed here)
# ═══════════════════════════════════════════════════════════════════

DIRECTIONS = ("forward", "backward", "left", "right", "360", "stop")
SPIN_DURATION_S = 10.0

# Keeps re-publishing the command so the motor controller watchdog
# (cmd_vel_timeout ~0.5s) does not stop the robot early, then sends stop.
_MOVE_SCRIPT = """
import sys, time
import rclpy
from std_msgs.msg import String

direction, duration_s = sys.argv[1], float(sys.argv[2])
rclpy.init()
node = rclpy.create_node('direct_motor_publisher')
pub = node.create_publisher(String, '/motor_command', 10)
deadline = time.time() + duration_s
try:
    while time.time() < deadline:
        pub.publish(String(data=direction))
        time.sleep(0.1)
finally:
    pub.publish(String(data='stop'))
    rclpy.shutdown()
"""


def _has_subscribers(topic: str) -> bool:
    try:
        result = subprocess.run(
            ["ros2", "topic", "info", topic],
            capture_output=True, text=True, timeout=3, check=False,
        )
    except Exception:
        return False
    match = re.search(r"Subscription count:\s*(\d+)", result.stdout)
    return result.returncode == 0 and bool(match) and int(match.group(1)) > 0


def _publish_once(topic: str, data: str) -> None:
    subprocess.Popen(
        ["ros2", "topic", "pub", "--once", topic, "std_msgs/msg/String", f"{{data: '{data}'}}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def move(direction: str, duration: float | None = None) -> str:
    """Drive in *direction* for *duration* seconds, then auto-stop.

    Default duration is 2s, or a full spin for "360".
    """
    direction = direction.strip().lower()
    if direction not in DIRECTIONS:
        return f"Unknown direction '{direction}'. Valid options: {', '.join(DIRECTIONS)}."

    if not _has_subscribers("/motor_command"):
        return (
            "Error: No subscribers on /motor_command. Start ROS first: "
            "ros2 launch gangubai_control motor_control.launch.py"
        )

    if direction == "stop":
        _publish_once("/motor_command", "stop")
        return "Robot stopped."

    if duration is None:
        duration = SPIN_DURATION_S if direction == "360" else 2.0
    if direction == "360":
        direction = "right"

    duration = max(0.2, float(duration))
    subprocess.Popen(
        ["/usr/bin/python3", "-c", _MOVE_SCRIPT, direction, str(duration)],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return f"Robot moving {direction} for {duration:.1f}s then auto-stopping."


def set_wander(action: str, check_subscribers: bool = True) -> str:
    """Start or stop autonomous wander mode."""
    action = action.strip().lower()
    if action not in ("start", "stop"):
        return f"Invalid wander action '{action}'. Use start or stop."

    if check_subscribers and not _has_subscribers("/wander_mode"):
        return "Error: No subscribers on /wander_mode. Start the ROS launch with the wander controller first."

    try:
        _publish_once("/wander_mode", action)
    except FileNotFoundError:
        return "Error: `ros2` command not found. Source ROS 2 first (source /opt/ros/humble/setup.bash)."

    if action == "start":
        return "Wander mode enabled. Robot will move autonomously until stopped."
    return "Wander mode stopped. Robot returned to idle."


# ═══════════════════════════════════════════════════════════════════
# RAG over the course slides
# ═══════════════════════════════════════════════════════════════════

_vectorstore = None


def _load_slide_chunks(file_path: str) -> list:
    """Parse one PPTX/PDF into overlapping chunks tagged with source + slide/page number."""
    import fitz  # PyMuPDF
    from langchain_core.documents import Document
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from pptx import Presentation

    file_name = os.path.basename(file_path)
    documents = []
    if file_path.lower().endswith(".pptx"):
        for i, slide in enumerate(Presentation(file_path).slides):
            text = "\n".join(s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text.strip())
            if text:
                documents.append(Document(page_content=text, metadata={"source": file_name, "type": "Slide", "number": i + 1}))
    else:
        for i, page in enumerate(fitz.open(file_path)):
            text = page.get_text().strip()
            if text:
                documents.append(Document(page_content=text, metadata={"source": file_name, "type": "Page", "number": i + 1}))

    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, add_start_index=True)
    chunks = splitter.split_documents(documents)
    # Put the citation info inside the text so the LLM can quote it.
    for chunk in chunks:
        m = chunk.metadata
        chunk.page_content = f"[{m['type']} {m['number']} from {m['source']}]\nContent: {chunk.page_content}"
    return chunks


def get_vectorstore():
    """Load the Chroma index from disk, or build it from SLIDES_DIR on first use."""
    global _vectorstore
    if _vectorstore is not None:
        return _vectorstore

    from langchain_community.vectorstores import Chroma
    from langchain_huggingface import HuggingFaceEmbeddings

    embeddings = HuggingFaceEmbeddings(model_name=config.EMBEDDING_MODEL)
    if os.path.exists(config.PERSIST_DIRECTORY):
        print("[RAG] Loading existing vector database...", flush=True)
        _vectorstore = Chroma(
            collection_name=config.CHROMA_COLLECTION,
            embedding_function=embeddings,
            persist_directory=config.PERSIST_DIRECTORY,
        )
        return _vectorstore

    print(f"[RAG] Building vector database from {config.SLIDES_DIR}...", flush=True)
    chunks = []
    for path in sorted(config.SLIDES_DIR.iterdir()):
        if path.suffix.lower() in (".pptx", ".pdf"):
            file_chunks = _load_slide_chunks(str(path))
            print(f"[RAG]   {path.name}: {len(file_chunks)} chunks", flush=True)
            chunks.extend(file_chunks)

    _vectorstore = Chroma.from_documents(
        documents=chunks,
        collection_name=config.CHROMA_COLLECTION,
        embedding=embeddings,
        persist_directory=config.PERSIST_DIRECTORY,
    )
    print(f"[RAG] Vector database created with {len(chunks)} chunks", flush=True)
    return _vectorstore


# ═══════════════════════════════════════════════════════════════════
# Tools
# ═══════════════════════════════════════════════════════════════════

@tool
def retrieve_context(query: str) -> str:
    """Retrieve information from the Gangubai knowledge base to help answer a query.

    Use this tool when you need to answer questions about topics in the knowledge base.

    Args:
        query: The question or topic to search for in the knowledge base

    Returns:
        Relevant information from the knowledge base with source citations
    """
    docs = get_vectorstore().similarity_search(query, k=3)
    return "\n\n".join(f"Source: {doc.metadata}\nContent: {doc.page_content}" for doc in docs)


@tool
def move_robot(
    direction: Literal["forward", "backward", "left", "right", "360", "stop"],
    duration: float | None = None,
) -> str:
    """Move the Gangubai robot in a given direction.

    Use this tool when the user asks the robot to move, drive, turn, spin,
    go forward/backward, rotate, or do a 360.

    Args:
        direction: The movement direction. One of:
            "forward"  – drive straight ahead
            "backward" – drive in reverse
            "left"     – turn/spin left
            "right"    – turn/spin right
            "360"      – do a full 360° clockwise spin
            "stop"     – immediately stop all motors
        duration: How long to move in seconds. Leave empty for the default
                  (2 seconds, or a full spin for "360"). Ignored for "stop".

    Returns:
        A status message confirming the action.
    """
    return move(direction, duration)


@tool
def set_wander_mode(action: Literal["start", "stop"]) -> str:
    """Start or stop autonomous wander behavior.

    Use this tool when the user asks the robot to wander/explore autonomously,
    or to stop wandering and return to idle.

    Args:
        action: "start" to enter wander mode, "stop" to return to idle.

    Returns:
        A status message confirming the mode change.
    """
    return set_wander(action)


@tool
def timer(duration_seconds: float) -> str:
    """Start a countdown timer overlay.

    Use this tool only when the user gave an explicit duration.
    If the duration is missing, ask the user for the time instead of calling the tool.

    Args:
        duration_seconds: Required countdown length in seconds.

    Returns:
        A confirmation message for the active timer overlay.
    """
    duration_seconds = max(1.0, float(duration_seconds))
    send_to_face("happy", event="timer_start", payload={"duration_seconds": duration_seconds})
    minutes, seconds = divmod(int(round(duration_seconds)), 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"Timer started for {hours}h {minutes:02d}m {seconds:02d}s."
    if minutes:
        return f"Timer started for {minutes}m {seconds:02d}s."
    return f"Timer started for {seconds}s."


@tool
def pomodoro(work_minutes: float = 25.0, break_minutes: float = 5.0, cycles: int = 4) -> str:
    """Start a pomodoro overlay.

    Use the standard 25 minute focus and 5 minute break split unless the user
    asks for different values. If the user does not specify cycles, use 4.

    Args:
        work_minutes: Focus length for each cycle.
        break_minutes: Break length between focus cycles.
        cycles: Number of focus cycles to run.

    Returns:
        A confirmation message for the active pomodoro overlay.
    """
    work_minutes = max(1.0, float(work_minutes))
    break_minutes = max(0.0, float(break_minutes))
    cycles = max(1, int(cycles))
    send_to_face("happy", event="pomodoro_start", payload={
        "work_seconds": work_minutes * 60.0,
        "break_seconds": break_minutes * 60.0,
        "cycles": cycles,
    })
    return (
        f"Pomodoro started with {cycles} cycles of {int(work_minutes)} minutes focus "
        f"and {int(break_minutes)} minutes break."
    )


# ═══════════════════════════════════════════════════════════════════
# Agent
# ═══════════════════════════════════════════════════════════════════

EMOTION_VALUES = Literal["happy", "sad", "angry", "curious", "excited", "confused", "neutral", "thinking"]
_EMOTION_DESCRIPTION = (
    "The single emotion that best matches this response. "
    "One of: happy, sad, angry, curious, excited, confused, neutral, thinking."
)


class ChatState(TypedDict):
    messages: Annotated[list[BaseMessage], add_messages]
    current_emotion: str


class ChatResponse(BaseModel):
    """Standard chatbot reply with an associated emotion."""
    content: str = Field(description="The chatbot's reply to the user.")
    emotion: EMOTION_VALUES = Field(description=_EMOTION_DESCRIPTION, default="neutral")


class RAGResponse(BaseModel):
    """RAG-grounded reply with citation and emotion."""
    explanation: str = Field(
        description="Detailed explanation answering the user's question based on the retrieved context"
    )
    citation: str = Field(
        description="Citation in the format: 'For more reference check <source>, <type> <number>' "
                    "e.g. 'For more reference check ECC_1_Introduction.pptx, Slide 3'"
    )
    emotion: EMOTION_VALUES = Field(description=_EMOTION_DESCRIPTION, default="neutral")


tools_list = [retrieve_context, move_robot, set_wander_mode, timer, pomodoro]
model = ChatGroq(model=config.GROQ_MODEL, temperature=config.TEMPERATURE)
llm_with_tools = model.bind_tools(tools_list)
llm_for_rag = model.with_structured_output(RAGResponse)
llm_for_chat = model.with_structured_output(ChatResponse)


# ── Keyword fallbacks (used when the LLM fails to emit a tool call) ──
_ROBOT_KEYWORDS = (
    "move", "turn", "rotate", "forward", "backward", "left", "right",
    "stop", "navigate", "motor", "robot",
)
_RETRIEVAL_CUES = (
    "teach", "explain", "detail", "what is", "how does", "why", "history",
    "evolution", "llm", "transformer", "nlp", "slides", "notes", "ppt",
    "chapter", "unit", "from the material", "according to the material",
)
_SMALLTALK_RE = re.compile(r"\b(joke|funny|humor|hello|hi|thanks|how are you|who are you|your name|thank you)\b")


def _looks_like_robot_action(query: str) -> bool:
    return any(keyword in query for keyword in _ROBOT_KEYWORDS)


def _should_force_retrieve(query: str) -> bool:
    if _looks_like_robot_action(query) or _SMALLTALK_RE.search(query):
        return False
    return any(cue in query for cue in _RETRIEVAL_CUES) and len(query.split()) >= 4


def _infer_direction(query: str) -> str:
    if any(token in query for token in ("360", "spin", "rotate", "turn around")):
        return "360"
    if "forward" in query or "ahead" in query:
        return "forward"
    if any(token in query for token in ("backward", "reverse", "back")):
        return "backward"
    for direction in ("left", "right", "stop"):
        if direction in query:
            return direction
    return "forward"


def _forced_tool_call(query: str) -> ChatState | None:
    """Return a forced tool-call state for movement or study questions, else None."""
    if _looks_like_robot_action(query):
        name, args, emotion = "move_robot", {"direction": _infer_direction(query)}, "excited"
    elif _should_force_retrieve(query):
        name, args, emotion = "retrieve_context", {"query": query}, "thinking"
    else:
        return None
    call = {"name": name, "args": args, "id": f"force_{name}_{uuid4().hex}", "type": "tool_call"}
    return {"messages": [AIMessage(content="", tool_calls=[call])], "current_emotion": emotion}


def _latest_user_query(messages: list[BaseMessage]) -> str:
    for msg in reversed(messages):
        if isinstance(msg, HumanMessage):
            return str(msg.content).strip()
    return ""


def _retrieved_context(messages: list[BaseMessage]) -> str | None:
    """Return retrieve_context output if it ran in the latest tool round-trip."""
    for msg in reversed(messages):
        if isinstance(msg, AIMessage):
            return None
        if isinstance(msg, ToolMessage) and msg.name == "retrieve_context":
            return str(msg.content)
    return None


def _answer_from_context(user_question: str, context: str) -> ChatState:
    # Clean prompt so the structured-output model doesn't see raw tool-call artefacts.
    rag_messages = [
        SystemMessage(content=config.SYSTEM_PROMPT),
        HumanMessage(
            content=(
                f"User question: {user_question}\n\n"
                f"Retrieved context:\n{context}\n\n"
                "Using the retrieved context, answer the user's question. "
                "Explain in your own words — simple and clear."
            )
        ),
    ]
    try:
        response: RAGResponse = llm_for_rag.invoke(rag_messages)
    except Exception:
        text = str(model.invoke(rag_messages).content).strip()
        return {
            "messages": [AIMessage(content=text or "I found relevant context, but had trouble formatting the response.")],
            "current_emotion": "thinking",
        }
    return {
        "messages": [AIMessage(content=f"{response.explanation}\n\n {response.citation}")],
        "current_emotion": response.emotion,
    }


def _plain_reply(all_messages: list[BaseMessage]) -> ChatState:
    try:
        response: ChatResponse = llm_for_chat.invoke(all_messages)
    except Exception:
        text = str(model.invoke(all_messages).content).strip()
        return {"messages": [AIMessage(content=text or "I am here.")], "current_emotion": "neutral"}
    return {"messages": [AIMessage(content=response.content)], "current_emotion": response.emotion}


def chat_node(state: ChatState) -> ChatState:
    messages = state["messages"]
    query = _latest_user_query(messages)

    context = _retrieved_context(messages)
    if context is not None:
        return _answer_from_context(query, context)

    all_messages = [SystemMessage(content=config.SYSTEM_PROMPT)] + list(messages)
    try:
        tool_response = llm_with_tools.invoke(all_messages)
        if tool_response.tool_calls:
            return {"messages": [tool_response]}
    except Exception:
        # Groq sometimes rejects its own tool-call output; fall through to keyword fallbacks.
        pass

    # Only force a tool before any tool has run this turn, otherwise we'd loop.
    if isinstance(messages[-1], HumanMessage):
        forced = _forced_tool_call(query.lower())
        if forced is not None:
            return forced
    return _plain_reply(all_messages)


graph = StateGraph(ChatState)
graph.add_node("chat", chat_node)
graph.add_node("tools", ToolNode(tools_list))
graph.add_edge(START, "chat")
graph.add_conditional_edges("chat", tools_condition)
graph.add_edge("tools", "chat")

# For LangGraph Studio (langgraph.json) – compiled without a checkpointer.
agent = graph.compile()
